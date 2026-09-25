from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from mini_rag.policy.models import ExtractedDocument

_DASHES = str.maketrans(
    {
        "\u2011": "-",
        "\u2013": "—",
        "\u2014": "—",
        "\u2212": "-",
        "\x97": "—",
    }
)


def extract_file(path: Path) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    else:
        raise ValueError(f"Unsupported policy file: {path.name}")
    return ExtractedDocument(source_file=path.name, text=normalize_text(text))


def load_policy_dir(path: Path) -> list[ExtractedDocument]:
    if not path.is_dir():
        raise ValueError(f"Policy directory was not found: {path}")
    files = sorted(
        child
        for child in path.iterdir()
        if child.is_file() and child.suffix.lower() in {".pdf", ".docx"}
    )
    if not files:
        raise ValueError(f"No PDF or Word files in {path}")
    return [extract_file(child) for child in files]


def normalize_text(text: str) -> str:
    lines = []
    for raw_line in text.translate(_DASHES).splitlines():
        line = re.sub(r"[ \t]+", " ", raw_line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    return "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )


def _extract_pdf(path: Path) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(path) as document:
        for page in document.pages:
            pages.append(_pdf_page_text(page))
    return "\n".join(pages)


def _pdf_page_text(page: Any) -> str:
    text = page.extract_text() or ""
    missing: list[str] = []
    for table in page.extract_tables() or []:
        for row in table:
            cells = [(cell or "").strip() for cell in row]
            cells = [cell for cell in cells if cell]
            if len(cells) < 2:
                continue
            if not all(cell in text for cell in cells):
                missing.append(" | ".join(cells))
    if not missing:
        return text
    return "\n".join(missing + [text])
